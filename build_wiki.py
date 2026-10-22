"""Genera las paginas de la WIKI (Obsidian) a partir de los Excel originales.

Uso:  python build_wiki.py

Regla de oro: NADA se inventa. Cada dato sale de una celda concreta de un Excel y cada pagina
dice de que libro/hoja/fila viene. Los enlaces [[...]] los crea este script (Obsidian no los
crea solo). No se importan PRECIOS ni descuentos: solo referencia, descripcion y litros.

Paginas generadas (carpetas regeneradas enteras en cada ejecucion; no editar a mano):
  productos/<marca>/   un producto = una referencia de una lista de precios
  familias_ad/         familias de aceite AD Parts del catalogo (homologaciones, refs)
  equivalencias/       una por fila de la matriz "Equivalencias Aceites por Marcas"
  normas/              ACEA, API, ILSAC, JASO, SAE, homologaciones OEM... (una por norma)
  oem/                 codigos OEM de recambio (Febi)
  index.md, lint.md    indice y problemas de datos detectados (incluye contradicciones entre fuentes)

Normas: se unifican solas las grafias que solo difieren en espacios/guiones/puntos/mayusculas
(el signo + SE CONSERVA: G12 != G12+ != G12++). Equivalencias entre nombres distintos las
confirma el usuario y viven en normas_alias.json.
"""
import json
import re
import shutil
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent
BASE = RAIZ / "BASE DE CONOCIMIENTO"
WIKI = RAIZ / "OBSIDIAN" / "Buscador Aceites" / "WIKI"
FUENTES_MD = RAIZ / "OBSIDIAN" / "Buscador Aceites" / "FUENTES_MD"
CARPETAS = ("productos", "familias_ad", "equivalencias", "normas", "oem")

# ---------------------------------------------------------------- utilidades
_alias_cfg = json.loads((RAIZ / "normas_alias.json").read_text(encoding="utf-8"))
ALIAS = {k: v for k, v in _alias_cfg.items() if not k.startswith("_")}


def limpio(v):
    if v is None:
        return ""
    if isinstance(v, float):
        v = int(v) if v.is_integer() else round(v, 3)
    return " ".join(str(v).replace("\xa0", " ").split())


def clave(s):
    # el + distingue normas (G12 / G12+ / G12++, LHM / LHM+): se conserva como PLUS
    return re.sub(r"[^A-Z0-9]", "", s.upper().replace("+", "PLUS"))


ALIAS_CLAVE = {clave(k): v for k, v in ALIAS.items()}
_ruta_dec = RAIZ / "decisiones.json"
_dec_json = json.loads(_ruta_dec.read_text(encoding="utf-8")) if _ruta_dec.exists() else {}
DECISIONES = _dec_json.get("decisiones", [])
CORR_EXCEL = _dec_json.get("correcciones_excel", [])


def fichero(s, largo=110):
    s = re.sub(r'[\\/:*?"<>|#^\[\]]', "-", s)
    return re.sub(r"\s+", " ", s).strip(" .")[:largo].strip(" .")


def nombre_nota(libro: Path, hoja: str):
    """Nombre de la nota Markdown de FUENTES_MD (misma regla que office_to_md.py)."""
    n = lambda s: re.sub(r'[\\/:*?"<>|]', "-", s).strip()
    return f"{n(libro.stem)} - {n(hoja)}"


def lista_plana(v):
    return [t.strip() for t in re.split(r"[;/|]", v) if t.strip()]


COMBOS = []  # combinaciones "A 504 00 / 507 00" partidas, para revisarlas en el lint


def lista_homologaciones(v):
    """Separa por ';'. Si un trozo es 'VW 504 00 / 507 00' (o '504.00/507.00') lo parte en
    'VW 504 00' y 'VW 507 00' heredando el prefijo de letras del primero."""
    out = []
    for parte in (t.strip() for t in v.split(";") if t.strip()):
        trozos = [q.strip() for q in parte.split("/")]
        m = re.search(r"\d", trozos[0])
        if len(trozos) > 1 and m and m.start() > 0 and all(re.fullmatch(r"[\d .\-]+", q) for q in trozos[1:]):
            pref = trozos[0][: m.start()]
            nuevos = [trozos[0]] + [pref + q for q in trozos[1:]]
            COMBOS.append((parte, nuevos))
            out += nuevos
        else:
            out.append(parte)
    return out


def mismos_litros(a, b):
    """La matriz escribe 200 L para el bidon (208/209/205 L en las listas): no es contradiccion."""
    if a is None or b is None:
        return True
    if abs(a - b) < 1e-6:
        return True
    return 200 <= a <= 210 and 200 <= b <= 210


def litros_num(v):
    try:
        return round(float(str(v).replace(",", ".")), 3)
    except ValueError:
        return None


RE_VISC = re.compile(r"(?<![0-9A-Za-z])(\d{1,2}W-?\d{2,3})(?![0-9])", re.I)


def visc_de(texto):
    m = RE_VISC.search(texto)
    return m.group(1).upper().replace(" ", "").replace("-", "") if m else ""


# ---------------------------------------------------------------- carga de fuentes
class Producto(dict):
    pass


PRODUCTOS = []
POR_REF = defaultdict(list)  # (marca, ref) -> [Producto]


def anadir_producto(marca, ref, desc, litros, fila, libro, hoja, extra=None):
    p = Producto(marca=marca, ref=ref, desc=desc, litros=litros, fila=fila,
                 libro=libro.name, nota=nombre_nota(libro, hoja), hoja=hoja,
                 normas=[], oem=[], extra=extra or {}, equiv=[], familia=[], febi=[], tcmatic=[])
    PRODUCTOS.append(p)
    POR_REF[(marca, ref)].append(p)
    return p


def cargar_listas_simples():
    fuentes = [
        ("AD Parts", "Productos AD Parts Abril 2026.xlsx"),
        ("AD Standard", "Productos AD Parts Estandard Abril 2026.xlsx"),
        ("Castrol", "Productos Castrol Abril 2026.xlsx"),
        ("Eni", "Productos Eni Abril 2026.xlsx"),
        ("Repsol", "Productos Repsol Abril 2026.xlsx"),
        ("Shell", "Productos Shell Abril 2026.xlsx"),
    ]
    for marca, nombre in fuentes:
        libro = BASE / "PRODUCTOS Y REFERENCIAS" / nombre
        ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
        for n, f in enumerate(ws.iter_rows(values_only=True), 1):
            ref, desc, lit = (limpio(f[i]) if i < len(f) else "" for i in range(3))
            if not ref or not desc or ref.upper() in ("REF", "REFERENCIA") or desc.upper() in ("PRODUCTO", "DESCRIPCION", "DESCRIPCIÓN"):
                continue
            p = anadir_producto(marca, ref, desc, lit, n, libro, ws.title)
            v = visc_de(desc)
            if v:
                p["normas"].append(("sae", f"SAE {v}"))
                p["visc_desc"] = v


def cargar_febi():
    libro = BASE / "PRODUCTOS Y REFERENCIAS" / "Productos Febi Abril 2026.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    col = dict(ref=1, articulo=2, oem=3, litros=4, acea=5, api=6, ilsac=7, approval=8, gear=9,
               hidraulico=10, refrigerante=11, dot=12, sae=13, color=16, uso=17)
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        if n == 1:
            continue
        d = {k: limpio(f[i]) if i < len(f) else "" for k, i in col.items()}
        if not d["ref"]:
            continue
        p = anadir_producto("febi", d["ref"], d["articulo"], d["litros"], n, libro, ws.title,
                            extra={"color": d["color"]})
        for campo, pref, tipo in (("acea", "ACEA", "acea"), ("api", "API", "api"), ("ilsac", "ILSAC", "ilsac"),
                                  ("sae", "SAE", "sae"), ("dot", "", "dot"), ("gear", "", "engranaje"),
                                  ("hidraulico", "", "hidraulico"), ("refrigerante", "Refrigerante", "refrigerante")):
            for t in lista_plana(d[campo]):
                p["normas"].append((tipo, f"{pref} {t}".strip()))
        for campo, tipo in (("approval", "homologacion"), ("uso", "uso")):
            for t in lista_homologaciones(d[campo]):
                p["normas"].append((tipo, t))
        for t in (x.strip() for x in d["oem"].split(";") if x.strip()):
            if re.search(r"\b(SAE|DOT|LTR|LITR|LITROS?)\b|^\d+W-?\d+", t, re.I):
                p["extra"]["oem_texto"] = t  # no es un codigo de recambio real
            else:
                p["oem"].append(t)
        v = visc_de(d["sae"])
        if v:
            p["visc_desc"] = v


def cargar_tcmatic():
    libro = BASE / "PRODUCTOS Y REFERENCIAS" / "Productos TC-Matic Abril 2026.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        if n == 1 or not f or not limpio(f[0]):
            continue
        g = lambda i: limpio(f[i]) if i < len(f) else ""
        p = anadir_producto("TC-Matic", g(0), g(1), g(2), n, libro, ws.title,
                            extra={"tipo_aceite": g(3), "tipo_caja": g(7)})
        for t in lista_plana(g(4)):
            p["normas"].append(("homologacion", t))


def cargar_refrigerantes_ad():
    libro = BASE / "PRODUCTOS Y REFERENCIAS" / "Productos AD Parts Refrigerantes Abril 2026.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        ref, desc = limpio(f[0]) if f else "", limpio(f[1]) if f and len(f) > 1 else ""
        if not ref.startswith("ADP") or not desc:
            continue
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*L\b", desc, re.I)  # la columna UNDS son unidades por palé, no litros
        anadir_producto("AD Refrigerantes", ref, desc, m.group(1).replace(",", ".") if m else "", n, libro, ws.title)


HOM_REPSOL = defaultdict(list)  # codigo RP (SIRDI) -> filas de "Repsol Homologaciones.xlsx"
CORRECCIONES_EXCEL = []  # correcciones de texto confirmadas por el usuario sobre el Excel de homologaciones (se listan en el lint)
MESES_N = {m: i for i, m in enumerate(["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                                       "septiembre", "octubre", "noviembre", "diciembre"], 1)}
COLS_NO_OEM = {"PAIS", "SIRDI", "NOMBRE ACTUAL", "NOMBRE", "NOMBRE ANTIGUO", "NOMBRE ANTERIOR", "SAE", "GRADO SAE", "ACEA", "API", "ILSAC", "JASO"}


def cargar_homologaciones_repsol():
    """Hojas de 'Repsol Homologaciones.xlsx' con columna SIRDI (= codigo RP_xxxxx del pie de las fichas).
    Se guardan SAE/ACEA/API/ILSAC/JASO y, aparte, cada columna de fabricante con su texto original."""
    import datetime
    libro = BASE / "HOMOLOGACIONES" / "Repsol" / "Repsol Homologaciones.xlsx"
    wb = openpyxl.load_workbook(libro, data_only=True)
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        h = None
        for i, fila in enumerate(filas[:8]):
            nombres = [limpio(c).upper() for c in fila]
            if "SIRDI" in nombres and "SAE" in nombres:
                h = i
                break
        if h is None:
            continue
        fecha_hoja = None  # la hoja lleva su fecha de actualizacion en las primeras filas
        for fila in filas[:4]:
            for c in fila:
                if isinstance(c, (datetime.datetime, datetime.date)):
                    fecha_hoja = (c.year, c.month)
        ix = {n: i for i, n in enumerate(nombres) if n}
        col = lambda *claves: next((ix[k] for k in claves if k in ix), None)
        for n, f in enumerate(filas[h + 1:], h + 2):
            g = lambda c: limpio(f[c]) if c is not None and c < len(f) else ""
            cod = g(col("SIRDI"))
            if not re.fullmatch(r"R[A-Z]_[0-9A-Z]{5}", cod):
                continue
            fila_d = dict(cod=cod, pais=g(col("PAIS")), nombre=g(col("NOMBRE ACTUAL", "NOMBRE")), antiguo=g(col("NOMBRE ANTIGUO", "NOMBRE ANTERIOR")),
                          sae=g(col("SAE", "GRADO SAE")), acea=g(col("ACEA")), api=g(col("API")), ilsac=g(col("ILSAC")), jaso=g(col("JASO")),
                          oem={nombres[i]: limpio(f[i]) for i in range(len(nombres)) if nombres[i] and nombres[i] not in COLS_NO_OEM
                               and i < len(f) and limpio(f[i]) and "FALTA SIRDI" not in limpio(f[i]).upper()},
                          hoja=ws.title, fila=n, fecha_hoja=fecha_hoja, nota=nombre_nota(libro, ws.title))
            HOM_REPSOL[cod].append(fila_d)


def corregir_excel_repsol(correcciones):
    """Aplica las correcciones de texto confirmadas por el usuario a las filas cargadas (el Excel original no se toca)."""
    for c in correcciones:
        n = 0
        for filas in HOM_REPSOL.values():
            for f in filas:
                if c["campo"] in f and c["de"] in f[c["campo"]]:
                    f[c["campo"]] = f[c["campo"]].replace(c["de"], c["a"])
                    n += 1
        CORRECCIONES_EXCEL.append((c, n))


RE_API = r"\b(?:S[A-Z]|C[A-Z]-?\d?|FA-4|GL-\d|PG-\d|T[A-D])(?:\s+PLUS)?\b"


def _sin_marcas(txt):
    return re.sub(r"(?i)en tr[aá]mite|\*|\([^)]*\)", " ", txt)


def tokens_api(txt):
    """Codigos API de un texto libre, normalizados con clave() para compararlos."""
    return {clave(x) for x in re.findall(RE_API, _sin_marcas(txt).upper())}


def tokens_api_lit(txt):
    return list(dict.fromkeys(re.findall(RE_API, _sin_marcas(txt).upper())))


def toks_norma(txt, quitar_edicion=True):
    """Lista de codigos de un campo tipo ACEA 'A3/B4-21' (el año de edicion se quita para enlazar: B4-21 -> B4)."""
    out = []
    for x in re.split(r"[/,;]", _sin_marcas(txt)):
        x = x.strip()
        if quitar_edicion:
            x = re.sub(r"-\d{2}$", "", x)
        if x:
            out.append(x)
    return list(dict.fromkeys(out))


def segmentos_hom(valor):
    """'229.5*,226.5*' / '504 00 / 507 00*en trámite' -> [{'valores': [...], 'formal': bool, 'tramite': bool}].
    Interpretacion (el Excel no trae leyenda): * = homologacion formal (como en las fichas), 'en trámite' = solicitada y aun no concedida.
    La marca del final de un segmento se aplica a todas sus alternativas separadas por '/'."""
    out = []
    for seg in [s for s in re.split(r"[;,]", valor) if s.strip()]:
        tramite = bool(re.search(r"en tr[aá]mite", seg, re.I))
        formal = "*" in seg
        limpio_s = re.sub(r"en tr[aá]mite", " ", seg, flags=re.I).replace("*", " ")
        limpio_s = re.sub(r"\(#\)|\(antes[^)]*\)", " ", limpio_s, flags=re.I)
        vals = [v.strip() for v in re.split(r"\s*/\s*", limpio_s) if v.strip()]
        if vals:
            out.append({"valores": vals, "formal": formal and not tramite, "tramite": tramite, "texto": seg.strip()})
    return out


EQUIV_B = []  # equivalencias "por bloques de marca": grasas, hidraulicos, motor vehiculo industrial, transmision manual y ejes
CFG_B = {  # libro -> (tipo, {columna de propiedad: significado}, unidad de la cantidad)
    "Equivalencias Grasas.xlsx": ("Grasas", {"COMP": "texto", "NLGI": "nlgi", "KG": "cant"}, "kg"),
    "Equivalencias Hidraulicos.xlsx": ("Hidráulicos", {"SAE": "sae", "DIN": "texto", "LITROS": "cant"}, "L"),
    "Equivalencias Motor Vehiculo Industrial.xlsx": ("Motor vehículo industrial", {"VISCO": "sae", "ACEA": "acea", "LITROS": "cant"}, "L"),
    "Equivalencias Transmisión Manual y Ejes.xlsx": ("Transmisión manual y ejes", {"SAE": "sae", "API": "api", "LITROS": "cant"}, "L"),
}
MARCA_B = {"REPSOL": "Repsol", "DYNAMIC": "Dynamic", "ENI": "Eni", "ADP": "AD Parts", "ADS": "AD Standard", "CASTROL": "Castrol", "SHELL": "Shell"}


def cargar_equiv_bloques():
    """Cada fila con >=2 marcas con referencia es una equivalencia. Las propiedades (SAE, NLGI...) se arrastran hacia abajo
    mientras la celda siguiente esta vacia; la cantidad (litros/kg) es de cada fila."""
    for nombre, (tipo, mapa, unidad) in CFG_B.items():
        libro = BASE / "EQUIVALENCIAS" / nombre
        ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
        filas = list(ws.iter_rows(values_only=True))
        marcas = [(i, MARCA_B.get(limpio(c).upper(), limpio(c))) for i, c in enumerate(filas[0]) if i > 0 and limpio(c)]
        nombres = [limpio(c).upper() for c in filas[1][:3]]
        actual = {}
        for n, f in enumerate(filas[2:], 3):
            g = lambda i: limpio(f[i]) if i < len(f) else ""
            if not any(g(i) for i in range(len(f))):
                continue
            if g(0):
                actual = {nombres[0]: g(0)}
            if g(1):
                actual[nombres[1]] = g(1)
            items = [(m, g(i), g(i + 1), g(i + 2)) for i, m in marcas if g(i)]
            if not items:
                continue
            props = {k: v for k, v in actual.items() if mapa.get(k) != "cant"}
            EQUIV_B.append(dict(tipo=tipo, props=props, mapa=mapa, cant=g(2), unidad=unidad, fila=n, libro=libro.name,
                                nota=nombre_nota(libro, ws.title), hoja=ws.title, items=items))


FAMILIAS = []  # catalogo AD Parts


def cargar_catalogo_ad():
    libro = BASE / "CATALOGOS" / "Ad Parts" / "AD PARTS - Catalogo Lubricantes 2025.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        if n == 1 or not limpio(f[2] if len(f) > 2 else ""):
            continue
        g = lambda i: limpio(f[i]) if i < len(f) else ""
        fam = dict(tipo=g(0), categoria=g(1), producto=g(2), visc=g(3), acea=g(4), api=g(5), ilsac=g(6),
                   jaso=g(7), homolog=g(8), refs=[r.strip() for r in g(9).split(";") if r.strip()],
                   fila=n, libro=libro.name, nota=nombre_nota(libro, ws.title), hoja=ws.title)
        fam["normas"] = []
        for campo, pref, tipo in (("acea", "ACEA", "acea"), ("api", "API", "api"), ("ilsac", "ILSAC", "ilsac"),
                                  ("jaso", "JASO", "jaso")):
            for t in lista_plana(fam[campo]):
                fam["normas"].append((tipo, f"{pref} {t}"))
        v = visc_de(fam["visc"])
        if v:
            fam["normas"].append(("sae", f"SAE {v}"))
        for t in lista_homologaciones(fam["homolog"]):
            fam["normas"].append(("homologacion", t))
        FAMILIAS.append(fam)


EQUIV = []  # matriz de equivalencias
CORR_MATRIZ = []  # filas de la matriz donde un valor API estaba en la columna ACEA
ESTADOS = ("SIN EQUIVALENCIA", "EN OTROS FORMATOS", "FUERA DE TARIFA", "SIN ACTUALIZAR", "SOLO EN 1 LITRO", "NO EN AD")
MARCAS_EQ = {8: "AD Parts", 9: "AD Standard", 10: "Eni", 11: "Shell", 12: "Castrol", 13: "Repsol"}


def cargar_equivalencias():
    libro = BASE / "EQUIVALENCIAS" / "Equivalencias Aceites por Marcas.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        if n == 1:
            continue
        g = lambda i: limpio(f[i]) if i < len(f) else ""
        if not any(g(i) for i in range(0, 5)):
            continue
        eq = dict(visc=g(0), acea=g(1), ilsac=g(2), ad=g(3), litros=g(4), fila=n, libro=libro.name,
                  nota=nombre_nota(libro, ws.title), hoja=ws.title, marcas={})
        for i, m in MARCAS_EQ.items():
            eq["marcas"][m] = g(i)
        # un valor API (GL-5 LS...) escrito en la columna ACEA es un error de la matriz: se pasa a API (confirmado por el usuario)
        # cada valor debe estar en la columna que le corresponde: un valor API (GL-5 LS) o una especificacion de ATF/hidraulico/cambio
        # (DEX II D, DSG, HLP2, DOT 5.1...) escrito en la columna ACEA NO es ACEA (confirmado por el usuario); se reclasifica y se anota
        eq["api"], eq["otras"], eq["tipo_cambio"] = [], [], []
        for campo, patron, columna in (("acea", r"[A-E]\d{1,2}(?:-\d{2})?|SIN ACEA", "ACEA"), ("ilsac", r"GF-\d[A-Z]?|SIN ILSAC", "ILSAC")):
            validos = []
            if campo == "acea":  # combinaciones ACEA con otros separadores: "C5 - A1", "C4 (C3)"
                cods = re.findall(r"[A-E]\d{1,2}(?:-\d{2})?", eq[campo])
                resto = re.sub(r"[\s\-/,;|()+]", "", re.sub(r"[A-E]\d{1,2}(?:-\d{2})?", "", eq[campo]))
                if cods and not resto:
                    eq[campo] = "; ".join(dict.fromkeys(cods))
                    continue
            for t in lista_plana(eq[campo]):
                if re.fullmatch(patron, t, re.I):
                    validos.append(t)
                elif re.match(r"(?i)^(GL|PG)-?\d|^F[AB]-?\d|^C[A-K]-?\d|^MT-?1$", t):
                    eq["api"].append(t)
                    CORR_MATRIZ.append(f"fila {n}: «{t}» en la columna {columna} → API")
                elif t.upper() in ("DSG", "DCT", "CVT"):  # tipo de caja de cambios, no es una norma (confirmado por el usuario)
                    eq["tipo_cambio"].append(t)
                    CORR_MATRIZ.append(f"fila {n}: «{t}» en la columna {columna} → tipo de caja de cambios (no es una norma)")
                else:
                    eq["otras"].append(t)
                    CORR_MATRIZ.append(f"fila {n}: «{t}» en la columna {columna} → especificación (no es {columna})")
            eq[campo] = "; ".join(validos)
        eq["api"] = "; ".join(eq["api"])
        for t in lista_plana(eq["visc"]):  # "ATF 2", "ATF +3 +4"...: normas de fluido, no viscosidades SAE
            if not re.fullmatch(r"\d{1,3}W?(-?\d{2,3})?", t):
                eq["otras"].append(t)
        EQUIV.append(eq)


FEBI_TC = []


def cargar_febi_tcmatic():
    libro = BASE / "EQUIVALENCIAS" / "Equivalencias entre Febi y TC-Matic.xlsx"
    ws = openpyxl.load_workbook(libro, data_only=True).worksheets[0]
    for n, f in enumerate(ws.iter_rows(values_only=True), 1):
        if n == 1 or not f or not limpio(f[0]):
            continue
        FEBI_TC.append(dict(tc=limpio(f[0]), febi=limpio(f[1]) if len(f) > 1 else "", fila=n,
                            libro=libro.name, nota=nombre_nota(libro, ws.title), hoja=ws.title))


# ---------------------------------------------------------------- construccion
def main():
    for c in (cargar_listas_simples, cargar_febi, cargar_tcmatic, cargar_refrigerantes_ad, cargar_catalogo_ad,
              cargar_homologaciones_repsol,
              cargar_equivalencias, cargar_equiv_bloques, cargar_febi_tcmatic):
        c()
    corregir_excel_repsol(CORR_EXCEL)

    # nombres de pagina de producto (unicos, sin colisiones ni siquiera por mayusculas)
    usados = set()
    cuenta = Counter((p["marca"], p["ref"]) for p in PRODUCTOS)
    for p in PRODUCTOS:
        base = fichero(f'{p["marca"]} {p["ref"]}')
        if cuenta[(p["marca"], p["ref"])] > 1:
            base = fichero(f'{p["marca"]} {p["ref"]} (fila {p["fila"]})')
        p["pagina"] = base
        usados.add(base.lower())

    # familias AD
    cf = Counter((f["producto"], f["visc"]) for f in FAMILIAS)
    for f in FAMILIAS:
        b = fichero(f'{f["producto"]} {f["visc"]}')
        f["pagina"] = b if cf[(f["producto"], f["visc"])] == 1 else fichero(f'{b} (fila {f["fila"]})')
    FAM_IDX = defaultdict(list)
    for f in FAMILIAS:
        FAM_IDX[(clave(f["producto"]), clave(f["visc"]))].append(f)

    # equivalencias: enlazar productos
    lint_eq = defaultdict(list)
    for eq in EQUIV:
        eq["pagina"] = fichero(f'Equiv {eq["visc"]} {eq["ad"] or "-"} {eq["litros"]} L (fila {eq["fila"]})')
        eq["enlaces"] = {}  # marca -> [Producto]
        eq["ad_fam"] = []
        if eq["ad"] and eq["ad"].upper() != "NO EN AD":
            eq["ad_fam"] = FAM_IDX.get((clave("AD " + eq["ad"]), clave(eq["visc"])), [])
            if not eq["ad_fam"]:  # familias sin viscosidad propia (ATF...): buscar solo por nombre
                eq["ad_fam"] = [f for f in FAMILIAS if clave(f["producto"]) == clave("AD " + eq["ad"])]
        if eq["ad"] and eq["ad"].upper() != "NO EN AD" and not eq["ad_fam"]:
            lint_eq["Familia AD de la matriz sin página en el catálogo AD Parts"].append(
                f'`AD {eq["ad"]}` (viscosidad/tipo `{eq["visc"]}`, equivalencia fila {eq["fila"]})')
        for marca, val in eq["marcas"].items():
            if not val or val.upper() in ESTADOS:
                continue
            prods = POR_REF.get((marca, val), [])
            if not prods:
                lint_eq[f"Referencia de la matriz que NO está en la lista de precios de {marca}"].append(
                    f'`{val}` (equivalencia fila {eq["fila"]})')
                continue
            eq["enlaces"][marca] = prods
            for p in prods:
                p["equiv"].append(eq)
            # --- comprobaciones cruzadas
            for p in prods:
                ln, le = litros_num(p["litros"]), litros_num(eq["litros"])
                if not mismos_litros(ln, le):
                    lint_eq["Litros distintos: lista de precios vs matriz de equivalencias"].append(
                        f'[[{p["pagina"]}]] lista={p["litros"]} L · matriz={eq["litros"]} L ([[{eq["pagina"]}]])')
                vd = p.get("visc_desc")
                if vd and eq["visc"] and clave(vd) != clave(eq["visc"]):
                    lint_eq["Viscosidad distinta: descripción del producto vs matriz"].append(
                        f'[[{p["pagina"]}]] «{p["desc"]}» → {vd} · matriz={eq["visc"]} ([[{eq["pagina"]}]])')
        # ACEA de la matriz vs catalogo
        for fam in eq["ad_fam"]:
            a1 = {clave(x) for x in lista_plana(eq["acea"]) if clave(x) != clave("SIN ACEA")}
            a2 = {clave(x) for x in lista_plana(fam["acea"])}
            if a1 and a2 and not (a1 <= a2 or a2 <= a1):  # subconjunto = no es contradiccion
                lint_eq["ACEA distinto: matriz vs catálogo AD Parts"].append(
                    f'[[{eq["pagina"]}]] matriz={eq["acea"]} · catálogo [[{fam["pagina"]}]]={fam["acea"]}')

    # equivalencias por bloques (grasas, hidraulicos, motor VI, transmision): enlazar el producto de cada marca o crearlo desde la propia tabla
    for eq in EQUIV_B:
        eq["pagina"] = fichero(f'Equiv {eq["tipo"]} ' + " ".join(eq["props"].values()) + f' {eq["cant"]} {eq["unidad"]} (fila {eq["fila"]})')
        eq["enlaces"], eq["normas"] = {}, []
        for k, v in eq["props"].items():
            clase = eq["mapa"].get(k)
            for t in lista_plana(v):
                if clase == "sae":
                    eq["normas"].append(("sae", "SAE " + re.sub(r"(?i)^SAE\s*", "", t).upper()))
                elif clase in ("acea", "api", "nlgi"):
                    eq["normas"].append((clase, f"{clase.upper()} {t}"))
        sae_prop = next((re.sub(r"(?i)^SAE\s*", "", v) for k, v in eq["props"].items() if eq["mapa"].get(k) == "sae"), "")
        for marca, ref, desc, cant in eq["items"]:
            prods = POR_REF.get((marca, ref), [])
            if not prods:
                p = anadir_producto(marca, ref, desc, cant or eq["cant"], eq["fila"], BASE / "EQUIVALENCIAS" / eq["libro"], eq["hoja"])
                p["pagina"] = fichero(f"{marca} {ref}")
                usados.add(p["pagina"].lower())
                v = visc_de(desc)
                if v:
                    p["visc_desc"] = v
                prods = [p]
                if marca != "Dynamic":  # Dynamic no tiene lista de precios en la base de conocimiento
                    lint_eq[f'Referencia de las tablas de equivalencias por bloques que NO está en la lista de precios de {marca}'].append(
                        f'`{ref}` ({eq["libro"]}, fila {eq["fila"]})')
            eq["enlaces"].setdefault(marca, []).extend(prods)
            for p in prods:
                p["equiv"].append(eq)
                if not mismos_litros(litros_num(p["litros"]), litros_num(cant or eq["cant"])):
                    lint_eq["Cantidad distinta: lista de precios vs tablas de equivalencias por bloques"].append(
                        f'[[{p["pagina"]}]] lista={p["litros"]} · tabla={cant or eq["cant"]} ([[{eq["pagina"]}]])')
                vd = p.get("visc_desc")
                if vd and sae_prop and re.fullmatch(r"\d{1,2}W-?\d{2,3}", sae_prop, re.I) and clave(vd) != clave(sae_prop):
                    lint_eq["Viscosidad distinta: descripción del producto vs tablas de equivalencias por bloques"].append(
                        f'[[{p["pagina"]}]] «{p["desc"]}» → {vd} · tabla={sae_prop} ([[{eq["pagina"]}]])')
        if len(eq["enlaces"]) < 2:
            lint_eq["Filas de las tablas por bloques con una sola marca (sin equivalente)"].append(f'[[{eq["pagina"]}]]')

    # familias AD <-> productos AD
    for fam in FAMILIAS:
        fam["prods"] = []
        for r in fam["refs"]:
            enc = POR_REF.get(("AD Parts", r)) or POR_REF.get(("AD Standard", r))
            if enc:
                fam["prods"] += enc
                for p in enc:
                    p["familia"].append(fam)
            else:
                lint_eq["Referencia del catálogo AD Parts que no está en ninguna lista de precios AD"].append(
                    f'`{r}` en [[{fam["pagina"]}]]')

    # fichas tecnicas PDF (FUENTES_MD) -> producto o familia, solo por coincidencia EXACTA de referencia
    bajo = {(m, r.lower()): ps for (m, r), ps in POR_REF.items()}
    fam_por_ref = defaultdict(list)
    for fam in FAMILIAS:
        for r in fam["refs"]:
            fam_por_ref[r.lower()].append(fam)
    REGLAS = {"AD PARTS/Aceites AD Normal": ("AD Parts", ""), "AD PARTS/Aceites AD Standard": ("AD Standard", ""),
              "AD PARTS/Refrigerantes": ("AD Refrigerantes", "ADP"), "REPSOL/Aceites Ligero": ("Repsol", ""),
              "AD PARTS/Grasas": (None, "")}
    raiz_fichas = FUENTES_MD / "FICHAS TÉCNICAS"
    fichas_sin = defaultdict(list)
    n_fichas = Counter()

    def lineas_ficha(md):
        t = md.read_text(encoding="utf-8").split("<!-- pagina 1 -->", 1)[-1]
        return [l.strip() for l in t.splitlines() if l.strip() and not l.startswith("|")]

    def nombres_ficha(md):
        """Claves de nombre de producto leidas del titulo de la ficha Repsol ('Lubricantes Vehículo Pesado GIANT 1010 SAE 50')."""
        claves = []
        for l in lineas_ficha(md)[:2]:
            m = re.search(r"\b([A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ0-9]{2,}\b.*)$", l)
            if m:
                claves.append(clave(m.group(1)))
        return claves

    def clave_nombre(desc):
        d = re.sub(r"\s+\d+(?:[.,]\d+)?\s*L\s*$", "", desc.replace("\xa0", " ").strip(), flags=re.I)
        return clave(d)

    idx_nombre = defaultdict(list)
    for p in PRODUCTOS:
        if p["marca"] == "Repsol":
            idx_nombre[clave_nombre(p["desc"])].append(p)

    # quimicos AD (limpiaparabrisas): sin lista de precios; la referencia y su formato salen del catalogo PDF
    cat_pdf = FUENTES_MD / "CATALOGOS" / "Ad Parts" / "AD PARTS - Catalogo Refrigerantes 2026.md"
    cat_txt = cat_pdf.read_text(encoding="utf-8") if cat_pdf.exists() else ""
    for md in sorted((raiz_fichas / "AD PARTS" / "Químicos").glob("*.md")):
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*L\s+" + re.escape(md.stem) + r"\b", cat_txt)
        if not m:
            continue
        ls = [re.sub(r"^Revisión:\s*\d{1,2}/\d{1,2}/\d{4}\s*", "", l) for l in lineas_ficha(md)]
        titulo_q = " ".join(" ".join(ls[:2]).split("DESCRIPCIÓN")[0].split())
        p = anadir_producto("AD Químicos", md.stem, titulo_q[:90], m.group(1).replace(",", "."), 0,
                            cat_pdf.with_suffix(".pdf"), "—")
        p["fuente_texto"] = (f'PDF «AD PARTS - Catalogo Refrigerantes 2026» (referencia y formato "{m.group(0)}") → '
                             f'[[CATALOGOS/Ad Parts/AD PARTS - Catalogo Refrigerantes 2026]]')
        p["nota"] = "CATALOGOS/Ad Parts/AD PARTS - Catalogo Refrigerantes 2026"
        p["pagina"] = fichero(f'AD Químicos {md.stem}')
        usados.add(p["pagina"].lower())
        bajo[("AD Químicos", md.stem.lower())] = [p]
    REGLAS["AD PARTS/Químicos"] = ("AD Químicos", "")

    desacuerdos_repsol = []
    for md in sorted(raiz_fichas.rglob("*.md")):
        sub = md.parent.relative_to(raiz_fichas).as_posix()
        enlace = f'FUENTES_MD/FICHAS TÉCNICAS/{sub}/{md.stem}'
        n_fichas[sub] += 1
        marca, pref = REGLAS.get(sub, (None, ""))
        destino = bajo.get((marca, (pref + md.stem).lower()), []) if marca else []
        if not destino:
            fams = fam_por_ref.get(md.stem.lower(), []) if sub.startswith("AD PARTS") else []
            for fam in fams:
                fam.setdefault("fichas", []).append(enlace)
            if not fams and sub.startswith("REPSOL"):
                # la ficha no se llama como una referencia de la lista: se busca (1) por el codigo RP_xxxxx del pie de la ficha,
                # que coincide con el principio de la referencia (validado: 20 de 20 fichas con referencia exacta), y si no (2) por NOMBRE
                codigos = re.findall(r"RP_[0-9A-Z]{5}", md.read_text(encoding="utf-8"))
                por_codigo = [p for p in PRODUCTOS if p["marca"] == "Repsol" and any(p["ref"][3:].upper().startswith(c[3:]) for c in codigos)]
                por_nombre = [p for c in nombres_ficha(md) for p in idx_nombre.get(c, [])]
                if por_codigo and por_nombre and {id(p) for p in por_codigo} != {id(p) for p in por_nombre}:
                    desacuerdos_repsol.append((enlace, por_codigo, por_nombre))
                elegidos, via = (por_codigo, "el código RP de la ficha") if por_codigo else (por_nombre, "el nombre del producto")
                vistos_p = set()
                for p in elegidos:
                    if id(p) not in vistos_p:
                        vistos_p.add(id(p))
                        p.setdefault("fichas", []).append(enlace)
                        p.setdefault("fichas_via", {})[enlace] = via
                if elegidos:
                    continue
            if not fams:
                fichas_sin[sub].append(md.stem)
        for p in destino:
            p.setdefault("fichas", []).append(enlace)

    # datos extraidos de las fichas AD por extract_fichas_ad.py (reglas, sin IA)
    ruta_datos = RAIZ / "fichas_ad_extraidas.json"
    datos_fichas = json.loads(ruta_datos.read_text(encoding="utf-8")) if ruta_datos.exists() else {}
    def decisiones_de(p):
        out = []
        for dec in DECISIONES:
            if dec.get("api_ficha_a_excel") or dec.get("api_ficha_manda"):
                continue  # estas se aplican (y se anotan) solo en los productos Repsol donde de verdad actuan
            f = dec.get("filtro", {})
            if f.get("marca") and f["marca"] != p["marca"]:
                continue
            if f.get("ref_regex") and not re.search(f["ref_regex"], p["ref"]):
                continue
            if f.get("desc_regex") and not re.search(f["desc_regex"], p["desc"], re.I):
                continue
            if f.get("familia") and not any(fa["pagina"] == f["familia"] for fa in p["familia"]):
                continue
            out.append(dec)
        return out

    for p in PRODUCTOS:
        p["decisiones"] = decisiones_de(p)
        for fi in p.get("fichas", []):
            clv = fi.split("FICHAS TÉCNICAS/AD PARTS/", 1)[-1]
            if clv in datos_fichas:
                d = datos_fichas[clv]
                for dec in p["decisiones"]:
                    for a, b in dec.get("api_sustituir", {}).items():
                        d["api"] = [b if x == a else x for x in d["api"]]
                p["ficha_datos"] = d
                p["ficha_enlace"] = fi
                p["normas_f"] = ([("acea", f"ACEA {x}") for x in d["acea"]] + [("api", f"API {x}") for x in d["api"]]
                                 + [("ilsac", f"ILSAC {x}") for x in d["ilsac"]] + [("jaso", f"JASO {x}") for x in d["jaso"]]
                                 + ([("sae", f'SAE {d["grado_sae"]}')] if d["grado_sae"] else [])
                                 + [("homologacion", x) for x in d["oem_ok"]])

    # datos extraidos de las fichas Repsol por extract_fichas_repsol.py (reglas, sin IA)
    ruta_r = RAIZ / "fichas_repsol_extraidas.json"
    datos_repsol = json.loads(ruta_r.read_text(encoding="utf-8")) if ruta_r.exists() else {}
    for p in PRODUCTOS:
        if p["marca"] != "Repsol":
            continue
        for fi in p.get("fichas", []):
            d = datos_repsol.get(fi.split("FICHAS TÉCNICAS/REPSOL/", 1)[-1])
            if not d:
                continue
            p.setdefault("fichas_repsol", []).append((fi, d))
            # decisiones confirmadas sobre el API de estas fichas frente a Repsol Homologaciones.xlsx
            filas_c = [h for c in d["codigos_rp"] for h in HOM_REPSOL.get(c, [])]
            api_ficha = lambda: tokens_api(" ".join(it["literal"][3:] for it in d["niveles"] if it["literal"].startswith("API")))
            for dec in DECISIONES:
                for a, b in dec.get("api_ficha_a_excel", {}).items():
                    if clave(a) in api_ficha() and any(clave(b) in tokens_api(h["api"]) for h in filas_c):
                        for it in d["niveles"]:
                            if it["literal"].startswith("API"):
                                it["literal"] = re.sub(r"\b" + re.escape(a) + r"\b", b, it["literal"])
                        if dec not in p["decisiones"]:
                            p["decisiones"].append(dec)
                for x in dec.get("api_ficha_manda", []):
                    if api_ficha() == {clave(x["ficha"])} and any(tokens_api(h["api"]) == {x["excel"]} for h in filas_c):
                        d["_api_manda"] = True
                        if dec not in p["decisiones"]:
                            p["decisiones"].append(dec)
            nf = p.setdefault("normas_f", [])
            formales = p.setdefault("formales_lit", [])
            for it in d["niveles"]:
                m = re.match(r"(ACEA|API|ILSAC|JASO)\s+(.*)", it["literal"])
                if m:
                    toks = [x.strip() for x in re.split(r"[,/;]", re.sub(r"\s*\([^)]*\)", "", m.group(2))) if x.strip()]
                    lits = [(m.group(1).lower(), f"{m.group(1)} {x}") for x in toks if x]
                else:
                    lits = [("homologacion", x) for x in lista_homologaciones(it["literal"])]
                nf += lits
                if it["formal"]:
                    formales += [l for _, l in lits]
            for g in d["props"].get("grado_sae", []):
                g = g.strip()
                if re.fullmatch(r"\d{1,2}W-?\d{2,3}", g):
                    nf.append(("sae", "SAE " + g.replace("-", "")))
                elif g.isdigit():
                    nf.append(("sae", f"SAE {g}"))

        # filas de Repsol Homologaciones.xlsx de este producto: el codigo SIRDI (RP_0059F) coincide con el principio de la referencia
        clave5 = p["ref"][3:8].upper()
        filas_h = [h for cod, hs in HOM_REPSOL.items() if cod[3:] == clave5 for h in hs]
        p["hom_filas"] = [h for h in filas_h if h["pais"].lower().startswith("espa")] or filas_h
        for h in p["hom_filas"]:
            sae_h = re.sub(r"(?i)^SAE\s*", "", h["sae"]).replace("*", "").strip()
            if re.fullmatch(r"\d{1,2}W-?\d{2,3}", sae_h):
                p["normas"].append(("sae", "SAE " + sae_h.replace("-", "")))
            h["segs"] = {}
            for campo, pref in (("acea", "ACEA"), ("api", "API"), ("ilsac", "ILSAC"), ("jaso", "JASO")):
                segs_e = segmentos_hom(h[campo])
                if segs_e:
                    h["segs"][pref] = segs_e
                for sg in segs_e:
                    txt = " / ".join(sg["valores"])
                    toks = tokens_api_lit(txt) if campo == "api" else toks_norma(txt)
                    # "en trámite" = tramitada: se enlaza aparte (normas_t), sin mezclarla con las cumplidas
                    p.setdefault("normas_t" if sg["tramite"] else "normas", []).extend((campo, f"{pref} {t}") for t in toks)
            h["segs"].update({col: segmentos_hom(v) for col, v in h["oem"].items()})
            for col, segs in h["segs"].items():
                if col not in h["oem"]:
                    continue  # las columnas ACEA/API/ILSAC/JASO ya se enlazaron arriba como su tipo de norma
                for sg in segs:
                    p.setdefault("normas_t" if sg["tramite"] else "normas", []).extend(("homologacion", f"{col} {v}") for v in sg["valores"])

    # decisiones confirmadas por el usuario: grado SAE corregido, notas y avisos que dejan de aplicar
    SUP_ACEA_PROD, SUP_ACEA_FAM, SUP_API_PROD, SUP_API_FAM = set(), set(), set(), set()
    for dec in DECISIONES:
        for nom in dec.get("familias", []):
            for fam in FAMILIAS:
                if fam["pagina"] == nom:
                    fam.setdefault("notas_decision", []).append(dec)
                    if dec.get("prevalece_ficha_acea"):
                        SUP_ACEA_FAM.add(nom)
                    if dec.get("api_union"):
                        fam["normas"] += [("api", f"API {x}") for x in dec["api_union"]]
                        SUP_API_FAM.add(nom)
    for p in PRODUCTOS:
        for dec in p["decisiones"]:
            if dec.get("api_union"):
                p["normas"] += [("api", f"API {x}") for x in dec["api_union"]]
                SUP_API_PROD.add(p["pagina"])
            if dec.get("visc"):
                p["normas"] = [(t, l) for t, l in p["normas"] if t != "sae"] + [("sae", f'SAE {dec["visc"]}')]
                p["visc_desc"] = dec["visc"]
            if dec.get("prevalece_ficha_acea"):
                SUP_ACEA_PROD.add(p["pagina"])

    # los avisos que una decision confirmada ya resuelve dejan de listarse (la decision queda registrada en el lint)
    con_visc = {p["pagina"] for p in PRODUCTOS if any(d.get("visc") for d in p["decisiones"])}
    clv_v = "Viscosidad distinta: descripción del producto vs matriz"
    clv_a = "ACEA distinto: matriz vs catálogo AD Parts"
    lint_eq[clv_v] = [x for x in lint_eq[clv_v] if not any(f"[[{n}]]" in x for n in con_visc)]
    lint_eq[clv_a] = [x for x in lint_eq[clv_a] if not any(f"[[{n}]]" in x for n in SUP_ACEA_FAM)]

    # febi <-> tcmatic
    for r in FEBI_TC:
        tcs, fbs = POR_REF.get(("TC-Matic", r["tc"]), []), POR_REF.get(("febi", r["febi"]), [])
        if not tcs:
            lint_eq["TC-Matic de la tabla Febi↔TC-Matic que no está en la lista TC-Matic"].append(f'`{r["tc"]}`')
        if not fbs:
            lint_eq["Febi de la tabla Febi↔TC-Matic que no está en la lista Febi"].append(f'`{r["febi"]}` (con TC-Matic `{r["tc"]}`)')
        for t in tcs:
            t["febi"] += fbs
        for fb in fbs:
            fb["tcmatic"] += tcs

    # ---------------- registro de normas
    # reglas de nombre confirmadas por el usuario (normas_alias.json)
    _sobrantes = [(re.compile(rf"^{re.escape(a)}\s+(?={re.escape(b)})", re.I)) for a, b in _alias_cfg.get("_prefijos_sobrantes", [])]
    _pref_eq = sorted(((re.compile(r"^(?:" + "|".join(re.escape(v) for v in sorted(g[1:], key=len, reverse=True)) + r")\s+", re.I), g[0])
                       for g in _alias_cfg.get("_prefijos_equivalentes", [])),
                      key=lambda x: -max(len(v) for g in _alias_cfg["_prefijos_equivalentes"] if g[0] == x[1] for v in g[1:]))
    _relleno = [(re.compile(rf"^{re.escape(m)}\s+{re.escape(w)}\s+", re.I), m) for m, w in _alias_cfg.get("_palabras_relleno", [])]
    _omitidas = [(re.compile(rf"^{re.escape(m)}\s+(?=\d)", re.I), f"{m} {w} ") for m, w in _alias_cfg.get("_palabras_omitidas", [])]
    _implica = {clave(k): v for k, v in _alias_cfg.get("_implica", {}).items()}
    log_reglas = Counter()

    def canon_lit(lit):
        lit = " ".join(lit.split())
        orig = lit
        for rx in _sobrantes:  # "VOLVO Mack EOS-4.5" -> "Mack EOS-4.5"; "VW G12+" -> "G12+"
            lit = rx.sub("", lit, count=1)
        if lit.count(")") > lit.count("("):  # parentesis de cierre suelto: "G12evo)"
            lit = lit.rstrip(")").rstrip()
        for rx, m in _relleno:
            lit = rx.sub(m + " ", lit)
        for rx, canon in _pref_eq:
            if lit.upper().startswith(canon.upper() + " "):
                break  # ya lleva el nombre canonico (evita "ARVIN MERITOR MERITOR")
            if rx.match(lit):
                lit = rx.sub(canon + " ", lit, count=1)
                break
        lit = re.sub(r"(?i)\b(\w+)(?:\s+\1\b)+", r"\1", lit)  # errata: palabra repetida seguida ("VALTRA VALTRA G2-08")
        for rx, nuevo in _omitidas:
            lit = rx.sub(nuevo, lit, count=1)
        if lit != orig:
            log_reglas[(orig, lit)] += 1
        return ALIAS_CLAVE.get(clave(lit), lit)

    def con_implicadas(pares):
        """[(tipo, literal)] -> mismos pares canonicos + las normas que cada uno incluye (GL-5 LS incluye GL-5)."""
        out = []
        for tipo, lit in pares:
            lit = canon_lit(lit)
            out.append((tipo, lit))
            out += [(tipo, canon_lit(x)) for x in _implica.get(clave(lit), [])]
        return out

    menciones = defaultdict(list)  # clave -> [(literal, tipo)]
    for origen in list(PRODUCTOS) + FAMILIAS:
        for destino, fuente_n in (("normas_n", "normas"), ("normas_f_n", "normas_f"), ("normas_t_n", "normas_t")):
            origen[destino] = []  # normas_f = segun la ficha; normas_t = tramitadas (en trámite en el Excel de Repsol)
            vistos = set()
            for tipo, lit in con_implicadas(origen.get(fuente_n, [])):
                k = clave(lit)
                if not k or k in vistos:
                    continue
                if destino == "normas_t_n" and any(k == kk for _, kk in origen["normas_n"] + origen["normas_f_n"]):
                    continue  # si ya consta como cumplida por otra fuente, no se repite como tramitada
                vistos.add(k)
                menciones[k].append((lit, tipo))
                origen[destino].append((tipo, k))
    for eq in EQUIV:
        eq["normas_n"] = []
        lista_eq = [("sae", f"SAE {t}") for t in lista_plana(eq["visc"]) if re.fullmatch(r"\d{1,3}W?(-?\d{2,3})?", t)]
        lista_eq += [("ilsac", f"ILSAC {t}") for t in lista_plana(eq["ilsac"]) if not t.upper().startswith("SIN")]
        lista_eq += [("acea", f"ACEA {t}") for t in lista_plana(eq["acea"]) if clave(t) != clave("SIN ACEA")]
        lista_eq += [("api", f"API {t}") for t in lista_plana(eq.get("api", ""))]
        lista_eq += [(("sae" if t.upper().startswith("SAE ") else "dot" if t.upper().startswith("DOT") else "homologacion"), t)
                     for t in eq.get("otras", [])]
        for tipo, lit in con_implicadas(lista_eq):
            menciones[clave(lit)].append((lit, tipo))
            eq["normas_n"].append((tipo, clave(lit)))

    for eq in EQUIV_B:
        eq["normas_n"] = []
        for tipo, lit in con_implicadas(eq["normas"]):
            menciones[clave(lit)].append((lit, tipo))
            eq["normas_n"].append((tipo, clave(lit)))

    # prefijo de fabricante "decorativo": el mismo codigo CON LETRAS con varios de estos fabricantes delante (API DCTF-2, VW DCTF-2)
    # es una sola norma, la del codigo suelto (confirmado por el usuario). Los codigos solo numericos (MB 229.5) NO se tocan.
    lit_de = {k: Counter(l for l, _ in lst).most_common(1)[0][0] for k, lst in menciones.items()}
    rx_decor = re.compile(r"^(" + "|".join(re.escape(x) for x in _alias_cfg.get("_prefijos_decorativos", [])) + r")[\s\-]+(.+)$", re.I)
    grupos_d = defaultdict(list)
    for k, lit in lit_de.items():
        m = rx_decor.match(lit)
        if m:
            grupos_d[clave(m.group(2))].append((k, clave(m.group(1)), m.group(2).strip()))
    KEYMAP, fusion_decor = {}, []
    for rk, items in grupos_d.items():
        # misma norma con >=2 fabricantes decorativos delante y un codigo con letras (los solo numericos, tipo "MB 229.5", no se tocan)
        if len(rk) >= 4 and len({pf for _, pf, _ in items}) >= 2 and re.search(r"[A-Za-z]", items[0][2]):
            bare = items[0][2]
            for k, _, _ in items:
                KEYMAP[k] = clave(bare)
            fusion_decor.append((bare, [lit_de[k] for k, _, _ in items]))
    NOM_FORZADO = {clave(b): b for b, _ in fusion_decor}  # el nombre de la pagina es el codigo suelto
    if KEYMAP:
        for k, nk in KEYMAP.items():
            if k != nk:
                menciones[nk] = menciones.get(nk, []) + menciones.pop(k)
        remap = lambda lst: list(dict.fromkeys((t, KEYMAP.get(k, k)) for t, k in lst))
        for origen in list(PRODUCTOS) + FAMILIAS:
            for destino in ("normas_n", "normas_f_n", "normas_t_n"):
                origen[destino] = remap(origen[destino])
        for eq in EQUIV + EQUIV_B:
            eq["normas_n"] = remap(eq["normas_n"])

    NOM, VARIANTES, CAT, usados_n = {}, {}, {}, set()
    for k, lst in sorted(menciones.items()):
        lits = Counter(l for l, _ in lst)
        nombre = fichero(lits.most_common(1)[0][0])
        if nombre.lower() in usados_n or nombre.lower() in usados:
            nombre = fichero(nombre + " (norma)")
        usados_n.add(nombre.lower())
        NOM[k] = fichero(NOM_FORZADO[k]) if k in NOM_FORZADO else nombre
        VARIANTES[k] = sorted(lits)
        CAT[k] = Counter(t for _, t in lst).most_common(1)[0][0]

    # ---------------- escribir
    for sub in CARPETAS:
        shutil.rmtree(WIKI / sub, ignore_errors=True)
        (WIKI / sub).mkdir(parents=True, exist_ok=True)
    (WIKI / "lint_febi.md").unlink(missing_ok=True)

    def vinc(k, lit=None):
        n = NOM[k]
        return f"[[{n}]]" if not lit or fichero(lit) == n else f"[[{n}|{lit}]]"

    def fuente(o):
        return f'Excel «{o["libro"]}», hoja «{o["hoja"]}», fila {o["fila"]} → [[{o["nota"]}]]'

    por_norma = defaultdict(lambda: defaultdict(list))
    res_tramite = []  # (pagina, fabricante, valor, codigo, texto) de cada homologacion "en trámite" del Excel de Repsol

    def fecha_ficha(d):
        m = re.match(r"(\w+)\s+(\d{4})", d.get("fecha", ""))
        return (int(m.group(2)), MESES_N.get(m.group(1).lower(), 0)) if m and m.group(1).lower() in MESES_N else None
    ETIQ = {"acea": "ACEA", "api": "API", "ilsac": "ILSAC", "jaso": "JASO", "sae": "Viscosidad SAE",
            "dot": "DOT", "engranaje": "Tipo de engranaje", "hidraulico": "Tipo hidráulico",
            "refrigerante": "Refrigerante", "homologacion": "Homologaciones", "uso": "Uso recomendado (Febi)", "nlgi": "NLGI"}

    def bullets_normas(o, clave_lista="normas_n"):
        g = defaultdict(list)
        for tipo, k in o[clave_lista]:
            g[tipo].append(vinc(k))
        return [f"- **{ETIQ[t]}:** " + " · ".join(g[t]) for t in ETIQ if g.get(t)]

    for p in PRODUCTOS:
        carpeta = WIKI / "productos" / fichero(p["marca"])
        carpeta.mkdir(parents=True, exist_ok=True)
        titulo = f'{p["marca"]} {p["ref"]} · {p["desc"]}'
        fm = ["---", "tipo: producto", f'marca: "{p["marca"]}"', f'referencia: "{p["ref"]}"',
              f'descripcion: "{p["desc"].replace(chr(34), chr(39))}"', f'litros: "{p["litros"]}"',
              f'fuente: "[[{p["nota"]}]]"'] + ([] if p.get("fuente_texto") else [f'fila_excel: {p["fila"]}']) + ["verificado: false", "---", ""]
        cu = [f"# {titulo}", ""]
        if p["litros"]:
            cu += [f'**Formato:** {p["litros"]} L', ""]
        for dec in p["decisiones"]:
            cu += [f'> ✅ **Decisión confirmada ({dec["fecha"]}):** {dec["texto"]}', ""]
        cu += bullets_normas(p)
        if p.get("normas_t_n"):
            cu.append("- **Tramitadas** (en trámite en Repsol Homologaciones.xlsx; pueden estar ya concedidas): "
                      + " · ".join(vinc(k) for _, k in p["normas_t_n"]))
        if p["marca"] not in ("febi", "TC-Matic") and p.get("visc_desc"):
            cu.append("  - *(viscosidad leída de la descripción del producto)*")
        if p["oem"]:
            cu.append("- **OEM de recambio:** " + " · ".join(f"[[{fichero('OEM ' + o)}]]" for o in p["oem"]))
        if p["extra"].get("oem_texto"):
            cu.append(f'- **OEM (texto, no es un código de recambio):** {p["extra"]["oem_texto"]}')
        for k, et in (("color", "Color"), ("tipo_aceite", "Tipo de aceite"), ("tipo_caja", "Tipo de caja")):
            if p["extra"].get(k):
                cu.append(f'- **{et}:** {p["extra"][k]}')
        if p["familia"]:
            cu.append("- **Familia AD Parts:** " + " · ".join(f'[[{f["pagina"]}]]' for f in p["familia"]))
        if p["febi"]:
            cu.append("- **Equivalente Febi:** " + " · ".join(f'[[{q["pagina"]}]]' for q in p["febi"]))
        if p["tcmatic"]:
            cu.append("- **Equivalente TC-Matic:** " + " · ".join(f'[[{q["pagina"]}]]' for q in p["tcmatic"]))
        for eq in p["equiv"]:
            otros = [f'[[{q["pagina"]}]] ({m})' for m, qs in eq["enlaces"].items() for q in qs if q is not p]
            cu.append(f'- **Equivalencias** ([[{eq["pagina"]}]]): ' + (" · ".join(otros) if otros else "sin más marcas con referencia"))
        for fi in p.get("fichas", []):
            nota_n = f' *(casada por {p["fichas_via"][fi]}, no por la referencia)*' if fi in p.get("fichas_via", {}) else ""
            cu.append(f"- **Ficha técnica (PDF convertido):** [[{fi}|{fi.rsplit('/', 1)[-1]}]]{nota_n}")
        if p.get("ficha_datos"):
            d = p["ficha_datos"]
            cu += ["", "## Datos de la ficha técnica", "",
                   f'*Leídos automáticamente de [[{p["ficha_enlace"]}|la ficha]] (reglas, sin IA). Sin verificar.*', ""]
            cu.append(f'- **Título:** {d["titulo"]}' + (f' · {d["tipo_producto"]}' if d["tipo_producto"] else ""))
            if d["formatos"]:
                cu.append("- **Formatos citados:** " + ", ".join(f"{x} L" for x in d["formatos"]))
            if d["fecha"]:
                cu.append(f'- **Fecha de la ficha:** {d["fecha"]}')
            cu += bullets_normas(p, "normas_f_n")
            if d["oem_ambiguo"]:
                cu.append("- **Filas de homologaciones no enlazadas** (celdas combinadas del PDF, revisar en la ficha): "
                          + " | ".join(d["oem_ambiguo"]))
        vistos_fr = set()
        for fi, d in p.get("fichas_repsol", []):
            huella = json.dumps([d["niveles"], d["tabla"], d["titulos"]], sort_keys=True)
            if huella in vistos_fr:  # varias fichas del mismo producto con idéntico contenido: se muestra una sola vez
                continue
            vistos_fr.add(huella)
            cu += ["", "## Datos de la ficha técnica", "",
                   f'*Leídos automáticamente de [[{fi}|la ficha]] (reglas, sin IA). Sin verificar.*', ""]
            cu.append("- **Producto en la ficha:** " + " · ".join(d["titulos"])
                      + (f' · {d["fecha"]}' if d["fecha"] else "") + (f' · código {", ".join(d["codigos_rp"])}' if d["codigos_rp"] else ""))
            cu += bullets_normas(p, "normas_f_n")
            ff = [vinc(clave(canon_lit(l))) for l in dict.fromkeys(p.get("formales_lit", []))]
            if ff:
                cu.append("- **Homologaciones formales** (marcadas con * en la ficha; el resto son recomendaciones): " + " · ".join(ff))
            if d["texto_libre"]:
                cu.append("- **Texto libre no enlazado:** " + " | ".join(d["texto_libre"]))
            if d["tabla"]:
                cu += ["", "| Propiedad | Unidad | Método | Valor |", "|---|---|---|---|"]
                cu += [f'| {f["propiedad"]} | {f["unidad"]} | {f["metodo"]} | {" / ".join(f["valores"])} |' for f in d["tabla"]]
        for h in p.get("hom_filas", []):
            cu += ["", "## Homologaciones según Repsol Homologaciones.xlsx", "",
                   f'*Hoja «{h["hoja"]}», fila {h["fila"]} ([[{h["nota"]}]]), {h["pais"] or "sin país"}'
                   + (f', hoja actualizada {h["fecha_hoja"][1]:02d}/{h["fecha_hoja"][0]}' if h["fecha_hoja"] else "")
                   + ". El Excel no trae leyenda: `*` se interpreta como homologación formal (igual que en las fichas) "
                     "y «en trámite» como TRAMITADA: solicitada, y puede estar ya concedida desde la fecha de la hoja.*", ""]
            if h["nombre"] or h["antiguo"]:
                cu.append(f'- **Nombre en el Excel:** {h["nombre"]}' + (f' (antes: {h["antiguo"]})' if h["antiguo"] else ""))
            cu.append("- **SAE / ACEA / API / ILSAC / JASO (texto del Excel):** "
                      + " · ".join(x for x in (h["sae"], h["acea"], h["api"], h["ilsac"], h["jaso"]) if x))
            todas_f = {clave(canon_lit(l)) for _, l in p.get("normas_f", [])}
            formales_f = {clave(canon_lit(l)) for l in p.get("formales_lit", [])}
            fechas_f = [x for x in (fecha_ficha(d) for _, d in p.get("fichas_repsol", [])) if x]
            f_max = max(fechas_f) if fechas_f else None
            filas_t = []
            for col, segs in h["segs"].items():
                for sg in segs:
                    estado = "tramitada (en trámite)" if sg["tramite"] else ("formal (*)" if sg["formal"] else "sin marca")
                    comp = ""
                    if sg["tramite"]:
                        if not p.get("fichas_repsol"):
                            cod_c, comp = "sin_ficha", "Sin ficha para contrastar"
                        elif not (f_max and h["fecha_hoja"] and f_max > h["fecha_hoja"]):
                            cod_c, comp = "sin_info", f'Sin información posterior: la ficha es de {f_max[1]:02d}/{f_max[0]}' if f_max else "Sin información posterior"
                        else:
                            ks = [clave(canon_lit(f"{col} {v}")) for v in sg["valores"]]
                            if all(k in formales_f for k in ks):
                                cod_c, comp = "confirmada", f"✅ La ficha de {f_max[1]:02d}/{f_max[0]} (posterior a la hoja) la da como homologación FORMAL"
                            elif any(k in todas_f for k in ks):
                                cod_c, comp = "recomendacion", f"La ficha de {f_max[1]:02d}/{f_max[0]} la lista SIN * (recomendación, no formal)"
                            else:
                                cod_c, comp = "no_mencionada", f"La ficha de {f_max[1]:02d}/{f_max[0]} (posterior) no la menciona: sigue en trámite o se retiró"
                        res_tramite.append((p["pagina"], col, " / ".join(sg["valores"]), cod_c, comp))
                    filas_t.append(f'| {col} | {" / ".join(sg["valores"])} | {estado} | {comp} |')
            if filas_t:
                cu += ["", "| Fabricante | Valor | Estado | Comprobación con la ficha |", "|---|---|---|---|"] + filas_t
        cu += ["", "## Fuente", "", p.get("fuente_texto") or fuente(p), ""]
        (carpeta / f'{p["pagina"]}.md').write_text("\n".join(fm + cu), encoding="utf-8")
        for k in dict.fromkeys(k for _, k in p["normas_n"] + p["normas_f_n"]):
            por_norma[k]["prod"].append(p)
        for k in dict.fromkeys(k for _, k in p["normas_t_n"]):
            por_norma[k]["tram"].append(p)

    for fam in FAMILIAS:
        fm = ["---", "tipo: familia_ad", f'producto: "{fam["producto"]}"', f'viscosidad: "{fam["visc"]}"',
              f'categoria: "{fam["tipo"]} / {fam["categoria"]}"', f'fuente: "[[{fam["nota"]}]]"',
              f'fila_excel: {fam["fila"]}', "verificado: false", "---", ""]
        cu = [f'# {fam["producto"]} {fam["visc"]}', "", f'{fam["tipo"]} · {fam["categoria"]}', ""]
        for dec in fam.get("notas_decision", []):
            cu += [f'> ✅ **Decisión confirmada ({dec["fecha"]}):** {dec["texto"]}', ""]
        cu += bullets_normas(fam)
        if fam["prods"]:
            cu += ["", "## Referencias", ""] + [f'- [[{q["pagina"]}]] · {q["litros"]} L' for q in fam["prods"]]
        if fam.get("fichas"):
            cu += ["", "## Fichas técnicas (PDF convertidos)", ""] + [f"- [[{fi}|{fi.rsplit('/', 1)[-1]}]]" for fi in fam["fichas"]]
        eqs = [e for e in EQUIV if fam in e["ad_fam"]]
        if eqs:
            cu += ["", "## Equivalencias con otras marcas", ""] + [f'- [[{e["pagina"]}]]' for e in eqs]
        cu += ["", "## Fuente", "", fuente(fam), ""]
        (WIKI / "familias_ad" / f'{fam["pagina"]}.md').write_text("\n".join(fm + cu), encoding="utf-8")
        for tipo, k in fam["normas_n"]:
            por_norma[k]["fam"].append(fam)

    for eq in EQUIV:
        fm = ["---", "tipo: equivalencia", f'viscosidad: "{eq["visc"]}"', f'familia_ad: "{eq["ad"]}"',
              f'litros: "{eq["litros"]}"', f'fuente: "[[{eq["nota"]}]]"', f'fila_excel: {eq["fila"]}',
              "verificado: false", "---", ""]
        cu = [f'# Equivalencia {eq["visc"]} · AD {eq["ad"] or "-"} · {eq["litros"]} L', ""]
        cu += bullets_normas(eq)
        if eq.get("tipo_cambio"):
            cu.append("- **Tipo de caja de cambios:** " + ", ".join(dict.fromkeys(eq["tipo_cambio"])) + " *(no es una norma)*")
        if eq["ad_fam"]:
            cu.append("- **Familia AD (catálogo):** " + " · ".join(f'[[{f["pagina"]}]]' for f in eq["ad_fam"]))
        cu += ["", "| Marca | Referencia |", "|---|---|"]
        for m, v in eq["marcas"].items():
            if m in eq["enlaces"]:
                cu.append(f"| {m} | " + " · ".join(f'[[{q["pagina"]}]]' for q in eq["enlaces"][m]) + " |")
            else:
                cu.append(f"| {m} | {v or '—'} |")
        cu += ["", "## Fuente", "", fuente(eq), ""]
        (WIKI / "equivalencias" / f'{eq["pagina"]}.md').write_text("\n".join(fm + cu), encoding="utf-8")
        for tipo, k in eq["normas_n"]:
            por_norma[k]["eq"].append(eq)

    for eq in EQUIV_B:
        fm = ["---", "tipo: equivalencia", f'categoria: "{eq["tipo"]}"'] + [f'{k.lower()}: "{v}"' for k, v in eq["props"].items()] + [
            f'cantidad: "{eq["cant"]} {eq["unidad"]}"', f'fuente: "[[{eq["nota"]}]]"', f'fila_excel: {eq["fila"]}', "verificado: false", "---", ""]
        cu = [f'# Equivalencia {eq["tipo"]} · ' + " · ".join(f"{k} {v}" for k, v in eq["props"].items()) + f' · {eq["cant"]} {eq["unidad"]}', ""]
        cu += bullets_normas(eq)
        cu += [f"- **{k}:** {v}" for k, v in eq["props"].items() if eq["mapa"].get(k) == "texto"]
        cu += ["", "| Marca | Referencia | Descripción en la tabla |", "|---|---|---|"]
        for marca, ref, desc, cant in eq["items"]:
            cu.append(f"| {marca} | " + " · ".join(f'[[{q["pagina"]}]]' for q in eq["enlaces"].get(marca, [])) + f" | {desc} |")
        cu += ["", "## Fuente", "", fuente(eq), ""]
        (WIKI / "equivalencias" / f'{eq["pagina"]}.md').write_text("\n".join(fm + cu), encoding="utf-8")
        for tipo, k in eq["normas_n"]:
            por_norma[k]["eq"].append(eq)

    for k, nombre in NOM.items():
        alias = [v for v in VARIANTES[k] if fichero(v) != nombre]
        fm = ["---", "tipo: norma", f"categoria: {CAT[k]}"]
        if alias:
            fm.append("aliases: [" + ", ".join(json.dumps(a, ensure_ascii=False) for a in alias) + "]")
        fm += ["---", ""]
        cu = [f"# {nombre}", ""]
        if alias:
            cu += ["Otras grafías encontradas en las fuentes: " + ", ".join(f"`{a}`" for a in alias), ""]
        d = por_norma[k]
        for et, lst, f in (("Familias AD Parts", d["fam"], lambda x: f'- [[{x["pagina"]}]]'),
                           ("Productos", d["prod"], lambda x: f'- [[{x["pagina"]}]] · {x["desc"]} · {x["litros"]} L'),
                           ("Productos con esta norma TRAMITADA (en trámite según Repsol; puede estar ya concedida)", d["tram"],
                            lambda x: f'- [[{x["pagina"]}]] · {x["desc"]} · {x["litros"]} L'),
                           ("Filas de la matriz de equivalencias", d["eq"], lambda x: f'- [[{x["pagina"]}]]')):
            if lst:
                cu += [f"## {et} ({len(lst)})", ""] + [f(x) for x in lst] + [""]
        (WIKI / "normas" / f"{nombre}.md").write_text("\n".join(fm + cu), encoding="utf-8")

    oem_idx = defaultdict(list)
    for p in PRODUCTOS:
        for o in p["oem"]:
            oem_idx[clave(o)].append((o, p))
    oem_comp = []
    for k, lst in oem_idx.items():
        o = Counter(x for x, _ in lst).most_common(1)[0][0]
        prods = sorted({p["pagina"]: p for _, p in lst}.values(), key=lambda q: q["pagina"])
        cu = [f"# OEM {o}", "", f"## Productos con este código ({len(prods)})", ""]
        cu += [f'- [[{q["pagina"]}]] · {q["desc"]} · {q["litros"]} L' for q in prods]
        refr = {tuple(sorted(NOM[kk] for t, kk in q["normas_n"] if t in ("refrigerante", "sae"))) for q in prods}
        if len({q["desc"] for q in prods}) > 1 or len(refr) > 1:
            cu += ["", "> ⚠ Código compartido por productos de distinto tipo, refrigerante o viscosidad. Revisar."]
            oem_comp.append((o, prods))
        (WIKI / "oem" / f'{fichero("OEM " + o)}.md').write_text(
            "\n".join(["---", "tipo: oem", "---", ""] + cu) + "\n", encoding="utf-8")

    # ---------------- lint
    L = ["---", "tipo: informe", "---", "", "# Lint de datos (generado automáticamente)", "",
         f"Generado: {date.today()} · Productos: {len(PRODUCTOS)} · Familias AD: {len(FAMILIAS)} · "
         f"Equivalencias: {len(EQUIV)} · Normas: {len(NOM)} · OEM: {len(oem_idx)}", "",
         "Cada punto es algo a revisar contra la fuente. Ninguno se ha corregido solo.", ""]

    def seccion(titulo, items, nota=""):
        L.extend(["", f"## {titulo} ({len(items)})", ""] + ([nota, ""] if nota else []) + [f"- {i}" for i in items])

    mult = {k: v for k, v in VARIANTES.items() if len({fichero(x) for x in v}) > 1}
    seccion("Normas con varias grafías unificadas", [f"{vinc(k)} ← " + " | ".join(f"`{x}`" for x in v) for k, v in sorted(mult.items())],
            "Se unifican solas (espacios, guiones, mayúsculas) o por `normas_alias.json` (confirmadas por Yakoba).")
    seccion("Combinaciones 'A / B' separadas en dos normas", [f"`{a}` → " + ", ".join(f"`{x}`" for x in b) for a, b in sorted(set((a, tuple(b)) for a, b in COMBOS))],
            "Regla: «VW 504 00 / 507 00» pasa a «VW 504 00» y «VW 507 00». Revisar que ninguna sea un código indivisible.")
    seccion("OEM compartido por productos distintos", [f"[[{'OEM ' + o}]]: " + ", ".join(f'[[{q["pagina"]}]]' for q in ps) for o, ps in oem_comp])
    dups = [f'{m} {r} (filas ' + ", ".join(str(p["fila"]) for p in POR_REF[(m, r)]) + ")" for (m, r), n in cuenta.items() if n > 1]
    seccion("Referencias repetidas dentro de una misma lista de precios", dups)
    for t in sorted(lint_eq):
        if lint_eq[t]:
            seccion(t, sorted(set(lint_eq[t])))
    seccion("Nombres de norma unificados por las reglas confirmadas (prefijos equivalentes, palabras de relleno, palabras omitidas)",
            [f"`{a}` → `{b}` ({n})" for (a, b), n in sorted(log_reglas.items())])
    seccion("Normas unificadas porque el prefijo de fabricante es decoración (código con letras; confirmado por el usuario)",
            [f"{' | '.join('`' + x + '`' for x in sorted(set(v)))} → [[{fichero(b)}]]" for b, v in sorted(fusion_decor)])
    if CORR_MATRIZ:
        seccion("Matriz de equivalencias: valores que no corresponden a su columna (reclasificados al leerlos)", CORR_MATRIZ)
    L += ["", f"## Decisiones confirmadas aplicadas ({len(DECISIONES)})", ""]
    for dec in DECISIONES:
        n = sum(1 for p in PRODUCTOS if dec in p["decisiones"])
        L.append(f'- **{dec["id"]}** ({dec["fecha"]}, {n} productos): {dec["texto"]}')
    for c, n in CORRECCIONES_EXCEL:
        L.append(f'- **{c["id"]}** ({c["fecha"]}, corrección de texto en {n} celdas de {c["libro"]}): {c["texto"]}')
    L += ["", "## Fichas técnicas PDF y su enlace con productos", ""]
    for sub in sorted(n_fichas):
        L.append(f"- {sub}: {n_fichas[sub] - len(fichas_sin[sub])} de {n_fichas[sub]} enlazadas (por referencia exacta o, en Repsol, por nombre de producto)")
    for sub in ("AD PARTS/Aceites AD Normal", "AD PARTS/Refrigerantes", "REPSOL/Aceites Ligero"):
        seccion(f"Fichas sin producto: {sub}", [f"`{x}`" for x in fichas_sin[sub]],
                "La ficha existe pero su nombre no coincide con ninguna referencia de las listas de abril ni del catálogo.")
    # ---- datos de la ficha tecnica frente a lista de precios y catalogo AD
    cmp = defaultdict(list)
    n_con_ficha = 0
    for p in PRODUCTOS:
        d = p.get("ficha_datos")
        if not d:
            continue
        n_con_ficha += 1
        v = f'[[{p["pagina"]}]]'
        if d["grado_sae"] and p.get("visc_desc") and clave(d["grado_sae"]) != clave(p["visc_desc"]):
            cmp["Ficha: grado SAE distinto del de la descripción del producto"].append(
                f'{v} ficha={d["grado_sae"]} · descripción «{p["desc"]}»')
        ln = litros_num(p["litros"])
        if d["formatos"] and ln is not None and not any(mismos_litros(ln, litros_num(x)) for x in d["formatos"]):
            cmp["Ficha: el formato del producto no figura entre los formatos de la ficha"].append(
                f'{v} producto={p["litros"]} L · ficha={", ".join(d["formatos"])} L')
        for fam in p["familia"]:
            fv = f'[[{fam["pagina"]}]]'
            for campo, et in (("acea", "ACEA"), ("api", "API")):
                if campo == "acea" and (p["pagina"] in SUP_ACEA_PROD or fam["pagina"] in SUP_ACEA_FAM):
                    continue  # decision confirmada: la ficha manda
                if campo == "api" and (p["pagina"] in SUP_API_PROD or fam["pagina"] in SUP_API_FAM):
                    continue  # decision confirmada: se fusionan los dos valores
                # en la comparacion ACEA se ignora el año de edicion ("B4-08" = "B4"); las paginas conservan la grafia original
                ed = (lambda x: re.sub(r"-\d{2}$", "", x)) if campo == "acea" else (lambda x: x)
                a1 = {clave(ed(x)) for x in d[campo]}
                a2 = {clave(ed(x)) for x in lista_plana(fam[campo])}
                if a1 and a2 and not (a1 <= a2 or a2 <= a1):
                    cmp[f"Ficha vs catálogo AD: {et} distinto"].append(
                        f'{v} ficha={"; ".join(d[campo])} · catálogo {fv}={fam[campo]}')
            cat_t = lista_homologaciones(fam["homolog"])
            solo_ficha = [x for x in d["oem_ok"] if not any(clave(x) == clave(t) for t in cat_t)]
            if solo_ficha and cat_t:
                cmp["Ficha vs catálogo AD: homologaciones que están en la ficha y NO en el catálogo"].append(
                    f'{v} · catálogo {fv}: ' + ", ".join(f"`{x}`" for x in solo_ficha))
            if not d["oem_ambiguo"]:
                solo_cat = [t for t in cat_t if not any(clave(t) == clave(x) for x in d["oem_ok"])]
                if solo_cat and d["oem_ok"]:
                    cmp["Ficha vs catálogo AD: homologaciones que están en el catálogo y NO en la ficha"].append(
                        f'{v} · catálogo {fv}: ' + ", ".join(f"`{x}`" for x in solo_cat))
        if d["oem_ambiguo"]:
            cmp["Fichas con filas de homologaciones ilegibles (celdas combinadas): revisar a mano"].append(
                f'{v} · ' + " | ".join(d["oem_ambiguo"])[:160])
    n_repsol = 0
    for p in PRODUCTOS:
        if not p.get("fichas_repsol"):
            continue
        n_repsol += 1
        v = f'[[{p["pagina"]}]]'
        grados = {clave(g.replace("-", "")) for _, d in p["fichas_repsol"] for g in d["props"].get("grado_sae", [])}
        if p.get("visc_desc") and grados and clave(p["visc_desc"]) not in grados:
            cmp["Ficha Repsol: grado SAE distinto del de la descripción del producto"].append(
                f'{v} descripción «{p["desc"]}» · ficha ' + ", ".join(sorted({g for _, d in p["fichas_repsol"] for g in d["props"].get("grado_sae", [])})))
        for fi, d in p["fichas_repsol"]:
            if d["texto_libre"]:
                cmp["Fichas Repsol con texto libre en las homologaciones (no enlazado, revisar)"].append(f'{v} · ' + " | ".join(d["texto_libre"])[:160])
    # ficha Repsol vs lista maestra "Repsol Homologaciones.xlsx" (clave exacta: codigo RP = columna SIRDI)
    def toks_hom(s, quitar_edicion=True):
        s = re.sub(r"en tr[aá]mite", " ", s, flags=re.I).replace("*", " ")
        s = re.sub(r"\([^)]*\)", " ", s)
        out = set()
        for x in re.split(r"[/,;]", s):
            x = x.strip()
            if quitar_edicion:
                x = re.sub(r"-\d{2}$", "", x)
            if x:
                out.add(clave(x))
        return out

    codigos_vistos, sin_fila_hom, n_cruzadas = set(), [], 0
    for p in PRODUCTOS:
        for fi, d in p.get("fichas_repsol", []):
            for cod in d["codigos_rp"]:
                if (cod, fi) in codigos_vistos:
                    continue
                codigos_vistos.add((cod, fi))
                filas = [h for h in HOM_REPSOL.get(cod, []) if h["pais"].lower().startswith("espa")] or HOM_REPSOL.get(cod, [])
                if not filas:
                    sin_fila_hom.append(f'`{cod}` (ficha `{fi.rsplit("/", 1)[-1]}`, p. ej. [[{p["pagina"]}]])')
                    continue
                n_cruzadas += 1
                h = filas[0]
                f_acea = {clave(re.sub(r"-\d{2}$", "", x.strip())) for it in d["niveles"] if it["literal"].startswith("ACEA")
                          for x in re.split(r"[,/;]", re.sub(r"\([^)]*\)", " ", it["literal"][4:])) if x.strip()}
                f_api = tokens_api(" ".join(it["literal"][3:] for it in d["niveles"] if it["literal"].startswith("API")))
                f_sae = {clave(g.replace("-", "")) for g in d["props"].get("grado_sae", [])}
                h_acea, h_api = toks_hom(h["acea"]), tokens_api(h["api"])
                sae_h = re.sub(r"(?i)^SAE\s*", "", h["sae"]).replace("*", "").strip()
                h_sae = {clave(sae_h.replace("-", ""))} if re.fullmatch(r"\d{1,2}W-?\d{2,3}|\d{2,3}", sae_h) else set()
                ref = f'`{cod}` [[{p["pagina"]}]] · Excel [[{h["nota"]}]] fila {h["fila"]}'
                for et, a, b, ta, tb in (("SAE", f_sae, h_sae, ", ".join(d["props"].get("grado_sae", [])), h["sae"]),
                                         ("ACEA", f_acea, h_acea, "; ".join(sorted(f_acea)), h["acea"]),
                                         ("API", f_api, h_api, "; ".join(sorted(f_api)), h["api"])):
                    if a and b and not (a <= b or b <= a) and not (et == "API" and d.get("_api_manda")):
                        cmp[f"Ficha Repsol vs Repsol Homologaciones.xlsx: {et} distinto"].append(f"{ref}: ficha={ta} · Excel={tb}")
    # posibles erratas de PALABRA: un nombre poco usado que solo difiere de otro muy usado en una palabra casi igual ("ALISON C4" / "ALLISON C-4").
    # Los codigos con otro numero o letra suelta (DTFR 12B110 / 15B110, MAN 324 NF / SNF) son normas distintas y no se marcan.
    import difflib
    uso = {k: len(v) for k, v in menciones.items()}
    por_tokens = defaultdict(list)
    for k, n in NOM.items():
        if uso.get(k, 0) >= 4:
            por_tokens[len(n.split())].append(n)
    erratas = []
    for k, n in NOM.items():
        if uso.get(k, 0) > 2:
            continue
        tk = n.split()
        for c in por_tokens.get(len(tk), []):
            ck = c.split()
            dif = [(x, y) for x, y in zip(tk, ck) if clave(x) != clave(y)]
            if len(dif) == 1 and dif[0][0].isalpha() and dif[0][1].isalpha() and min(len(dif[0][0]), len(dif[0][1])) >= 3                     and difflib.SequenceMatcher(None, dif[0][0].upper(), dif[0][1].upper()).ratio() >= 0.8 and clave(n) != clave(c):
                erratas.append(f"[[{n}]] ({uso.get(k, 0)}) ← ¿[[{c}]]?")
                break
    if erratas:
        cmp["Posibles erratas de palabra en nombres de norma (parecen otra norma muy usada; revisar la fuente)"] = sorted(erratas)
    # posibles alias: mismo codigo/valor pero distinto nombre de fabricante entre fuentes ("MB 229.71" / "MB-Approval 229.71")
    grupos_v = defaultdict(set)
    for k, nombre in NOM.items():
        m = re.match(r"^((?:[^\W\d_]+[\s\-\.]+)+?)(\S*\d.*)$", nombre)
        if m:
            grupos_v[clave(m.group(2))].add((m.group(1).strip(" -."), nombre))
    cand_alias = []
    for rk, vs in grupos_v.items():
        if len(rk) >= 4 and len({clave(pf) for pf, _ in vs}) > 1:
            cand_alias.append(" ↔ ".join(f"[[{n}]]" for _, n in sorted(vs)))
    if cand_alias:
        cmp["Posibles alias entre fuentes: mismo código con distinto nombre de fabricante (confirmar cuáles son lo mismo)"] = sorted(cand_alias)
    usados_h = {h["cod"] for p in PRODUCTOS for h in p.get("hom_filas", [])}
    huerf = [f'`{h["cod"]}` {h["nombre"]} · [[{h["nota"]}]] fila {h["fila"]}' for hs in HOM_REPSOL.values() for h in hs
             if h["pais"].lower().startswith("espa") and h["cod"] not in usados_h and h["hoja"] in ("PCMO", "HD", "MOTO 4T", "MOTO 2T", "MTFs", "ATFs", "OOPP", "AGRO")]
    if huerf:
        cmp["Productos de automoción en Repsol Homologaciones.xlsx (España) sin ninguna referencia en la lista de precios de abril"] = huerf
    if sin_fila_hom:
        cmp[f"Fichas Repsol cuyo código RP no está en Repsol Homologaciones.xlsx (cruzadas con éxito: {n_cruzadas})"] = sin_fila_hom
    for enlace_f, pc, pn in desacuerdos_repsol:
        cmp["Ficha Repsol: el código RP y el nombre apuntan a productos distintos (se usó el código)"].append(
            f'`{enlace_f.rsplit("/", 1)[-1]}` código→' + ", ".join(f'[[{x["pagina"]}]]' for x in pc[:3])
            + " · nombre→" + ", ".join(f'[[{x["pagina"]}]]' for x in pn[:3]))
    for cod_c, tit in (("confirmada", "Repsol «en trámite» que una ficha POSTERIOR da ya como homologación formal"),
                       ("recomendacion", "Repsol «en trámite» que la ficha posterior lista sin * (recomendación, no formal)"),
                       ("no_mencionada", "Repsol «en trámite» que la ficha posterior no menciona (sigue en trámite o se retiró)")):
        lst = sorted({f"[[{pg}]] · {col} {val}" for pg, col, val, c, _ in res_tramite if c == cod_c})
        if lst:
            cmp[tit] = lst
    sin_info_t = sorted({f"[[{pg}]] · {col} {val} ({txt})" for pg, col, val, c, txt in res_tramite if c in ("sin_info", "sin_ficha")})
    if sin_info_t:
        cmp["Repsol «en trámite» sin ficha posterior con la que contrastar (comprobar en la web de Repsol)"] = sin_info_t
    cuentas_t = Counter(c for _, _, _, c, _ in res_tramite)
    L += ["", f"## Datos de fichas leídos por reglas: AD Parts {n_con_ficha} productos · Repsol {n_repsol} productos", "",
          f"- Homologaciones Repsol «en trámite» (filas de producto × fabricante): {len(res_tramite)} → "
          f"confirmadas por ficha posterior: {cuentas_t['confirmada']} · listadas sin * en la ficha: {cuentas_t['recomendacion']} · "
          f"no mencionadas en la ficha posterior: {cuentas_t['no_mencionada']} · sin ficha posterior con la que contrastar: "
          f"{cuentas_t['sin_info'] + cuentas_t['sin_ficha']}", ""]
    for t in sorted(cmp):
        seccion(t, list(dict.fromkeys(cmp[t])))
    sin_lit = [f'[[{p["pagina"]}]]' for p in PRODUCTOS if not p["litros"]]
    seccion("Productos sin litros en la lista", sin_lit)
    en_eq = {(m, p["ref"]) for e in EQUIV for m, ps in e["enlaces"].items() for p in ps}
    L += ["", "## Cobertura de la matriz de equivalencias", ""]
    for m in MARCAS_EQ.values():
        tot = sum(1 for p in PRODUCTOS if p["marca"] == m)
        L.append(f"- {m}: {sum(1 for p in PRODUCTOS if p['marca'] == m and (m, p['ref']) in en_eq)} de {tot} productos aparecen en la matriz")
    (WIKI / "lint.md").write_text("\n".join(L) + "\n", encoding="utf-8")

    # ---------------- indice
    I = ["---", "tipo: indice", "---", "", "# Índice de la wiki", "",
         "Generado por `build_wiki.py`. Fuentes sin tocar en `FUENTES_MD`.", "", "- [[lint]] — problemas de datos y contradicciones entre fuentes", "",
         "## Productos por marca", ""]
    for m, n in sorted(Counter(p["marca"] for p in PRODUCTOS).items()):
        I.append(f"- {m}: {n}")
    I += ["", f"## Familias AD Parts ({len(FAMILIAS)})", ""] + [f'- [[{f["pagina"]}]]' for f in sorted(FAMILIAS, key=lambda f: f["pagina"])]
    I += ["", f"## Equivalencias entre marcas, aceites de motor ({len(EQUIV)})", ""] + [f'- [[{e["pagina"]}]]' for e in EQUIV]
    for tp in dict.fromkeys(e["tipo"] for e in EQUIV_B):
        I += ["", f"## Equivalencias entre marcas, {tp}", ""] + [f'- [[{e["pagina"]}]]' for e in EQUIV_B if e["tipo"] == tp]
    (WIKI / "index.md").write_text("\n".join(I) + "\n", encoding="utf-8")

    print(f"productos={len(PRODUCTOS)} familias={len(FAMILIAS)} equivalencias={len(EQUIV)} normas={len(NOM)} oem={len(oem_idx)}")
    for t in sorted(lint_eq):
        print(f"  lint: {t}: {len(set(lint_eq[t]))}")
    print(f"  combos partidos={len(set((a, tuple(b)) for a, b in COMBOS))} grafias_unificadas={len(mult)} oem_compartidos={len(oem_comp)}")


if __name__ == "__main__":
    main()
